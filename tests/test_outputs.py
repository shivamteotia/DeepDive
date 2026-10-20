import csv
import io
import json

import openpyxl

import outputs
from tests.test_agent import Script, call, chat, say
from tests.test_api_extraction import add, extract, screened, with_extraction

HEADERS = ["Title", "Year"]
ROWS = [["=HYPERLINK(\"http://x\")", 2020], ["+1 trick", None], ["Plain, with \"quotes\"", 2019]]


def test_csv_opens_in_excel_and_neutralises_formulas():
    data = outputs.to_csv(HEADERS, ROWS)
    assert data.startswith("﻿".encode())  # Excel needs the BOM to read UTF-8
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == HEADERS and rows[1][0].startswith("'=") and rows[2] == ["'+1 trick", ""]
    assert rows[3] == ["Plain, with \"quotes\"", "2019"]


def test_xlsx_round_trips_without_formulas():
    wb = openpyxl.load_workbook(io.BytesIO(outputs.to_xlsx(HEADERS, ROWS, "Included")))
    ws = wb.active
    assert ws.title == "Included" and [c.value for c in ws[1]] == HEADERS
    assert ws["A2"].data_type == "s" and ws["A2"].value.startswith("'=") and ws["B2"].value == 2020


def _export(api, h, pid, **params):
    return api.get(f"/api/v1/projects/{pid}/export", headers=h, params=params)


def test_export_has_the_tabs_papers_decisions_and_extracted_columns(api, signup):
    h, pid, llm = screened(api, signup)
    with_extraction(api, llm)
    extract(api, h, pid, add(api, h, pid).json()["id"])
    r = _export(api, h, pid, tab="included", format="csv")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert r.headers["content-disposition"].endswith('.csv"')
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert len(rows) == 10 and {x["Decision"] for x in rows} == {"include"}
    assert {x["Study design"] for x in rows} == {"retrospective cohort"}
    assert all(x["Study design (quote)"] and x["Study design (source)"] == "abstract" for x in rows)
    wb = openpyxl.load_workbook(io.BytesIO(_export(api, h, pid, tab="included", format="xlsx").content))
    assert wb.active.max_row == 11 and isinstance(wb.active["C2"].value, int)  # year is a number


def test_export_is_project_scoped_and_validated(api, signup):
    h, pid, _ = screened(api, signup)
    other = signup("b@example.com")
    assert _export(api, other, pid, tab="all", format="csv").status_code == 404
    assert _export(api, h, pid, tab="duplicates", format="csv").status_code == 422
    assert _export(api, h, pid, tab="all", format="pdf").status_code == 422


def test_the_agent_offers_the_export_as_a_download(api, signup):
    h, pid, _ = screened(api, signup)
    script = Script([call("export_table", {"tab": "included", "format": "xlsx"})], say("Here is your table."))
    out = chat(api, h, pid, script, "export the included papers to Excel")[1].json()
    tool = next(m for m in out["messages"] if m["role"] == "tool")
    assert tool["content"]["result"]["export"] == {"tab": "included", "format": "xlsx", "papers": 10,
                                                   "columns": tool["content"]["result"]["export"]["columns"]}
    assert "Decision" in tool["content"]["result"]["export"]["columns"]
    assert json.dumps(tool["content"]["result"]).count("http") == 0  # a path for the UI, never an external link


def test_xlsx_keeps_numbers_like_text_and_only_guards_formulas():
    ws = openpyxl.load_workbook(io.BytesIO(outputs.to_xlsx(["v"], [["-0.5"], ["@x"], ["=1+1"]], "S"))).active
    assert [ws[f"A{i}"].value for i in (2, 3)] == ["-0.5", "@x"]
    assert ws["A4"].data_type == "s" and not str(ws["A4"].value).startswith("=")


def test_export_works_past_the_database_parameter_limit(api, signup, monkeypatch):
    import funnel
    h, pid, llm = screened(api, signup)
    with_extraction(api, llm)
    extract(api, h, pid, add(api, h, pid).json()["id"])
    sizes, real = [], funnel._decisions

    def spy(s, version, ids, test=False):
        sizes.append(len(ids))
        return real(s, version, ids, test)
    monkeypatch.setattr(funnel, "_decisions", spy)
    monkeypatch.setattr(funnel, "EXPORT_CHUNK", 3)
    rows = list(csv.DictReader(io.StringIO(_export(api, h, pid, tab="included", format="csv").content.decode("utf-8-sig"))))
    assert max(sizes) <= 3 and len(rows) == 10 and {x["Study design"] for x in rows} == {"retrospective cohort"}
